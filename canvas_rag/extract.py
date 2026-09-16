from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path


class _TextExtractor(HTMLParser):
    BLOCK = {"p", "div", "section", "article", "h1", "h2", "h3", "h4", "li", "tr", "br", "td", "th", "pre", "blockquote"}
    DROP = {"script", "style", "noscript"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.drop_depth = 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if self.drop_depth:
            if tag in self.DROP:
                self.drop_depth += 1
        elif tag in self.DROP:
            self.drop_depth = 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self.drop_depth:
            if tag in self.DROP:
                self.drop_depth -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.drop_depth:
            self.parts.append(data)


def html_to_markdown(source: str) -> str:
    parser = _TextExtractor()
    parser.feed(source)
    text = "".join(parser.parts).replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(line for line in lines if line)).strip()


def extract_attachment_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
        return html_to_markdown(raw) if suffix in {".html", ".htm"} else raw
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation
        slides = Presentation(str(path)).slides
        return "\n\n".join("\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)) for slide in slides)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        return "\n".join(
            f"[{sheet.title}] " + " | ".join(str(value) for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
            for sheet in book.worksheets
        )
    return ""
